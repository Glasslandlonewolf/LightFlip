"""Conservative paragraph reconstruction shared by editable Office exports.

Geometry, panel color and native PDF line metadata prevent cross-column
merges. Each result is one text box containing one paragraph and soft breaks,
not a group of independent line objects. Ambiguous lines stay independent.
"""
import math
import re
from statistics import median


def _height(row):
    return max(1., row['bbox'][3]-row['bbox'][1])


def _flow_height(row):
    # A final word without capitals/descenders has a shorter ink box, although
    # it uses the same text line. Compare font frames rather than ink alone.
    return max(_height(row), row['size']*.70)


def _color_distance(a, b):
    return math.sqrt(sum((((a >> shift) & 255)-((b >> shift) & 255))**2
                         for shift in (0, 8, 16)))


def _list_start(text):
    return bool(re.match(r'^\s*(?:[•●▪◆\-–]|\d{1,2}[.)、])\s+', text))


def _compatible(a, b, graphics):
    ax, ay, ar, ab = a['bbox']
    bx, by, br, bb = b['bbox']
    h = median((_flow_height(a), _flow_height(b)))
    if abs(a.get('angle', 0)) > 3 or abs(b.get('angle', 0)) > 3:
        return False
    # PDF extraction blocks are geometric, not semantic paragraphs: an Office
    # PDF or a loosely spaced paragraph may put every line in a separate block.
    # A block transition with additional vertical space remains a boundary.
    if 'block_id' in a and a['block_id'] != b.get('block_id') and by-ab > h*.55:
        return False
    if not .72 < _flow_height(b)/_flow_height(a) < 1.39:
        return False
    if not .30*h < by-ay < 2.15*h or not -.12*h <= by-ab <= 1.12*h:
        return False
    if min(len(a['text'].strip()), len(b['text'].strip())) < 4 or _list_start(b['text']):
        return False
    if _color_distance(a['color'], b['color']) > 35:
        return False
    if a.get('panel_color') is not None and b.get('panel_color') is not None:
        if _color_distance(a['panel_color'], b['panel_color']) > 25:
            return False
    left = abs(ax-bx) < h*.85
    centre = abs((ax+ar-bx-br)/2) < h*.55
    if not left and not centre:
        return False
    # A short bold label above a longer body is a heading, even if its OCR
    # height happens to resemble the body. Do not join it to that body.
    if len(a['text']) < 25 and len(b['text']) > len(a['text'])*1.5 and a['flags'] & 16:
        return False
    if 'width_ratio' not in a and abs(a['size']-b['size']) > min(a['size'], b['size'])*.18:
        return False
    x0, x1 = max(ax, bx), min(ar, br)
    mid = (ab+by)/2
    for graphic in graphics:
        gx, gy, gr, gb = graphic['bbox']
        if gx <= min(ax,bx) and gr >= max(ar,br) and gy <= ay and gb >= bb:
            continue  # A containing panel is a background, not a separator.
        barrier = gy <= mid <= gb or (graphic.get('geometry')=='line' and ab <= gy <= by)
        if max(0, min(x1, gr)-max(x0, gx)) > min(ar-ax, br-bx)*.35 and barrier:
            return False
    return True


def _uniform_ocr_style(rows):
    """Choose one installed font and stable weight for a raster paragraph."""
    from PIL import ImageFont
    from ocr_layout import fonts
    choices = fonts()
    scores = {}
    for row in rows:
        scores[row['font']] = scores.get(row['font'], 0)+len(row['text'])*row.get('font_fit', .65)
    family = max(scores, key=scores.get)
    regular_fraction = sum(not row['flags'] & 16 for row in rows)/len(rows)
    bold = regular_fraction < .30
    italic = all(row['flags'] & 2 for row in rows)
    path = next((path for name, path, b, i in choices if (name, b, i) == (family, bold, italic)), None)
    if path:
        # Glyph height is more stable than unrelated substituted font sizes.
        base = ImageFont.truetype(path, 100)
        estimates = [100*_height(row)/max(1, base.getbbox(row['text'])[3]-base.getbbox(row['text'])[1])
                     for row in rows]
        size = median(estimates)
    else:
        size = median(row['size'] for row in rows)
    for row in rows:
        row.update(font=family, size=size, flags=(16 if bold else 0)|(2 if italic else 0))
        if path:
            font = ImageFont.truetype(path, max(5, round(size)))
            box = font.getbbox(row['text'])
            # Office wraps according to advance width, not the visible ink box.
            row['width_ratio'] = font.getlength(row['text'])*1.025/max(1, row['bbox'][2]-row['bbox'][0])
            ascent, _ = font.getmetrics()
            row['origin'] = (row['bbox'][0], row['bbox'][1]-(box[1]-ascent))


def _merge_inline(rows):
    """Keep PDF inline emphasis inside one visual line, using native line IDs."""
    grouped = {}
    result = []
    for row in rows:
        if 'line_id' not in row:
            result.append(row)
        else:
            grouped.setdefault((row['block_id'], row['line_id']), []).append(row)
    for spans in grouped.values():
        spans.sort(key=lambda row:row['bbox'][0])
        first = dict(spans[0])
        if len(spans) > 1:
            runs, text = [], ''
            for span in spans:
                word = span['text']
                if runs and span['bbox'][0]-runs[-1]['bbox'][2] > first['size']*.12:
                    word = ' '+word
                runs.append(dict(span, text=word))
                text += word
            first.update(text=text, runs=runs,
                         bbox=(min(r['bbox'][0] for r in spans), min(r['bbox'][1] for r in spans),
                               max(r['bbox'][2] for r in spans), max(r['bbox'][3] for r in spans)))
        result.append(first)
    return result


def group_paragraphs(objects, page_width, page_height):
    lines = _merge_inline([dict(obj) for obj in objects if obj['kind'] == 'text'])
    graphics = [obj for obj in objects if obj['kind'] != 'text' and
                (obj['bbox'][2]-obj['bbox'][0])*(obj['bbox'][3]-obj['bbox'][1]) < page_width*page_height*.80]
    remaining = sorted(lines, key=lambda row:(row['bbox'][1], row['bbox'][0]))
    result = [obj for obj in objects if obj['kind'] != 'text']
    while remaining:
        rows = [remaining.pop(0)]
        while True:
            candidates = [row for row in remaining if _compatible(rows[-1], row, graphics)]
            if not candidates:
                break
            next_row = min(candidates, key=lambda row:(row['bbox'][1], abs(row['bbox'][0]-rows[-1]['bbox'][0])))
            # The closest matching line above belongs to this paragraph. This
            # prevents a wide caption from stealing a neighbouring column's row.
            competing = [row for row in remaining if row is not next_row and _compatible(row, next_row, graphics)]
            if any(row['bbox'][1] > rows[-1]['bbox'][1]+_height(rows[-1])*.35 for row in competing):
                break
            remaining.remove(next_row)
            rows.append(next_row)
        if len(rows) == 1:
            result.append(rows[0])
            continue
        if all('width_ratio' in row for row in rows):
            _uniform_ocr_style(rows)
        first = dict(rows[0])
        starts = [row['origin'][1]-row['size']*.92 for row in rows]
        step = median(starts[i+1]-starts[i] for i in range(len(starts)-1))
        # OCR lines occasionally fit a bad ascender. Baseline-independent glyph
        # steps bound the spacing before a single Office paragraph is created.
        if step <= 0:
            step = median(rows[i+1]['bbox'][1]-rows[i]['bbox'][1] for i in range(len(rows)-1))
        lefts = [row['bbox'][0] for row in rows]
        centres = [(row['bbox'][0]+row['bbox'][2])/2 for row in rows]
        centred = max(centres)-min(centres) < max(lefts)-min(lefts) and max(lefts)-min(lefts) > first['size']*.35
        first.update(paragraph_rows=rows, line_spacing=step, alignment='center' if centred else 'left',
                     text='\n'.join(row['text'] for row in rows),
                     bbox=(min(lefts), min(r['bbox'][1] for r in rows), max(r['bbox'][2] for r in rows), max(r['bbox'][3] for r in rows)),
                     origin=(min(lefts), rows[0]['origin'][1]), seq=min(row.get('seq', 0) for row in rows))
        result.append(first)
    return sorted(result, key=lambda obj:obj.get('seq', 0))


def add_textbox(slide, obj, scale=1., ox=0., oy=0.):
    """One native DrawingML paragraph; breaks preserve the scanned line layout."""
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    x0, y0, x1, y1 = obj['bbox']
    size = obj['size']*scale
    top = obj['origin'][1]-obj['size']*.92
    rows = obj.get('paragraph_rows', [obj])
    height = max(y1-y0+obj['size']*.65, obj.get('line_spacing', obj['size'])*len(rows)+obj['size']*.65)
    shape = slide.shapes.add_textbox(Pt(ox+x0*scale), Pt(oy+top*scale),
                                   Pt((x1-x0+obj['size']*.40)*scale), Pt(height*scale))
    shape.name = 'OCR 可编辑文段' if 'width_ratio' in obj and len(rows)>1 else '可编辑文字'
    shape.rotation = obj.get('angle', 0)
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = bool(obj.get('paragraph_rows'))
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    p = frame.paragraphs[0]
    p.space_before = p.space_after = Pt(0)
    if 'line_spacing' in obj:
        p.line_spacing = Pt(obj['line_spacing']*scale)
    p.alignment = PP_ALIGN.CENTER if obj.get('alignment') == 'center' else PP_ALIGN.LEFT
    for index, row in enumerate(rows):
        if index:
            p._p.append(OxmlElement('a:br'))
        for piece in row.get('runs', [row]):
            run = p.add_run()
            run.text = ''.join(c for c in piece['text'] if c in '\t\n' or ord(c)>=32)
            run.font.name = piece['font']
            run.font.size = Pt(max(1, piece['size']*scale))
            run.font.bold = bool(piece['flags'] & 16)
            run.font.italic = bool(piece['flags'] & 2)
            c = piece['color']
            run.font.color.rgb = RGBColor((c>>16)&255, (c>>8)&255, c&255)
            props = run._r.get_or_add_rPr()
            for tag in ('a:ea', 'a:cs'):
                node = OxmlElement(tag)
                node.set('typeface', piece['font'])
                props.append(node)
            if 'width_ratio' in piece:
                spacing = (piece['bbox'][2]-piece['bbox'][0])*(1-piece['width_ratio'])/max(1,len(piece['text'])-1)
                props.set('spc', str(round(max(-piece['size']*.08,min(piece['size']*.25,spacing))*scale*100)))
    return shape
