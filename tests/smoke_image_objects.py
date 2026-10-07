"""Integration acceptance for independently editable images (Windows runtime).

Run in the prepared LightFlip main Python environment. --runtime points to a
directory containing the existing OCR runtime. Fixtures are generated locally;
no user document, network, or third-party photograph is required.
"""
from pathlib import Path
import argparse
import io
import json
import random
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageChops, ImageStat
import pymupdf as fitz
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from core import convert
from pdf_ocr import image_pdf

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
}


def texture(seed, size=(270, 285)):
    """A self-created photographic texture, distinct from flat page decoration."""
    rng = random.Random(seed)
    width, height = size
    values = bytearray()
    for y in range(height):
        for x in range(width):
            noise = rng.randint(-28, 28)
            values.extend((max(0, min(255, 60 + x // 2 + noise)),
                           max(0, min(255, 45 + y // 2 + noise)),
                           max(0, min(255, 80 + (x + y) // 4 + noise))))
    image = Image.frombytes('RGB', size, bytes(values)).filter(ImageFilter.GaussianBlur(.55))
    draw = ImageDraw.Draw(image)
    for _ in range(60):
        x, y = rng.randrange(width), rng.randrange(height)
        draw.line((x, y, x + rng.randrange(-30, 31), y + rng.randrange(-30, 31)),
                  fill=tuple(rng.randrange(45, 210) for _ in range(3)), width=2)
    draw.ellipse((width*.25, height*.12, width*.73, height*.62), fill=(194, 158, 129))
    draw.polygon([(width*.47, height*.52), (width*.12, height), (width*.91, height)],
                 fill=(58+seed*3, 66+seed*2, 73+seed))
    draw.ellipse((width*.38, height*.30, width*.41, height*.33), fill=(42, 36, 31))
    draw.ellipse((width*.57, height*.30, width*.60, height*.33), fill=(42, 36, 31))
    return image


def fixture(folder):
    size = (1200, 760)
    base = Image.new('RGB', size, (246, 246, 243))
    ImageDraw.Draw(base).rectangle((0, 425, 1199, 759), fill=(52, 63, 78))
    source = base.copy()
    draw = ImageDraw.Draw(source)
    draw.text((55, 35), '轻转图片独立编辑',
              font=ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc', 40), fill=(32, 46, 58))
    draw.text((55, 104), 'Editable photo objects',
              font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 32), fill=(32, 46, 58))
    boxes = [(65, 200, 335, 485), (460, 215, 730, 500), (855, 190, 1125, 475)]
    for index, box in enumerate(boxes):
        source.paste(texture(index + 1), box[:2])
    source.save(folder/'mixed-image.png')
    bare = base.copy()
    for index, box in enumerate(boxes):
        bare.paste(texture(index + 1), box[:2])
    bare.save(folder/'photo-only.png')
    texture(7, size).save(folder/'full-photo.png')
    image_pdf(folder/'mixed-image.png', folder/'mixed-scan.pdf')
    image_pdf(folder/'photo-only.png', folder/'photo-only-scan.pdf')
    base.save(folder/'expected-background.png')
    return boxes, size, base


def all_text(path):
    with zipfile.ZipFile(path) as archive:
        names = [name for name in archive.namelist() if name == 'word/document.xml'
                 or name.startswith('ppt/slides/slide') and name.endswith('.xml')]
        return '\n'.join(''.join(node.text or '' for node in ET.fromstring(archive.read(name)).iter()
                                 if node.tag.endswith('}t')) for name in names)


def picture_objects(path):
    """Return visible photo objects and the single separately stored background."""
    if path.suffix == '.pptx':
        presentation = Presentation(path)
        assert len(presentation.slides) == 1
        pw, ph = presentation.slide_width, presentation.slide_height
        objects = []
        for shape in presentation.slides[0].shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            objects.append({'bbox': [shape.left/pw, shape.top/ph,
                                     (shape.left+shape.width)/pw, (shape.top+shape.height)/ph],
                            'name': shape.name, 'data': shape.image.blob})
    else:
        objects = []
        with zipfile.ZipFile(path) as archive:
            document = ET.fromstring(archive.read('word/document.xml'))
            sect = document.find('.//w:sectPr', NS)
            page_size = sect.find('w:pgSz', NS)
            pw = int(page_size.get('{'+NS['w']+'}w'))*635
            ph = int(page_size.get('{'+NS['w']+'}h'))*635
            rels = ET.fromstring(archive.read('word/_rels/document.xml.rels'))
            mapping = {node.get('Id'): node.get('Target') for node in rels}
            for anchor in document.findall('.//wp:anchor', NS):
                blip = anchor.find('.//a:blip', NS)
                if blip is None:
                    continue
                posx = int(anchor.find('wp:positionH/wp:posOffset', NS).text)
                posy = int(anchor.find('wp:positionV/wp:posOffset', NS).text)
                extent = anchor.find('wp:extent', NS)
                ex, ey = int(extent.get('cx')), int(extent.get('cy'))
                target = mapping[blip.get('{'+NS['r']+'}embed')]
                objects.append({'bbox': [posx/pw, posy/ph, (posx+ex)/pw, (posy+ey)/ph],
                                'name': anchor.find('wp:docPr', NS).get('name'),
                                'data': archive.read('word/'+target)})
    backgrounds = [obj for obj in objects if
                   (obj['bbox'][2]-obj['bbox'][0])*(obj['bbox'][3]-obj['bbox'][1]) > .8]
    assert len(backgrounds) == 1, (path.name, 'one background required', len(backgrounds))
    return [obj for obj in objects if obj not in backgrounds], backgrounds[0]


def iou(a, b):
    inter = max(0, min(a[2], b[2])-max(a[0], b[0])) * max(0, min(a[3], b[3])-max(a[1], b[1]))
    union = (a[2]-a[0])*(a[3]-a[1]) + (b[2]-b[0])*(b[3]-b[1]) - inter
    return inter/max(union, 1e-9)


def acceptance(path, boxes, size, base, needs_text=True):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith('.xml') or name.endswith('.rels'):
                ET.fromstring(archive.read(name))
    objects, background = picture_objects(path)
    expected = [[x0/size[0], y0/size[1], x1/size[0], y1/size[1]] for x0, y0, x1, y1 in boxes]
    matches = []
    for box in expected:
        assert objects, (path.name, 'no independent image objects')
        best = max(objects, key=lambda obj: iou(obj['bbox'], box))
        score = iou(best['bbox'], box)
        assert score > .84, (path.name, 'missing or merged photograph', box, best['bbox'], score)
        matches.append(best)
    assert len({tuple(obj['bbox']) for obj in matches}) == len(boxes), 'multiple photos merged'
    assert len({obj['data'] for obj in matches}) == len(boxes), 'photos reused the same flattened media'
    cleaned = Image.open(io.BytesIO(background['data'])).convert('RGB').resize(size)
    for box in boxes:
        # The centre must have been removed from the base image, rather than
        # merely covered by another copy that exposes a ghost when moved.
        x0, y0, x1, y1 = box
        inner = (x0+25, y0+25, x1-25, y1-25)
        difference = ImageChops.difference(cleaned.crop(inner), base.crop(inner))
        error = sum(ImageStat.Stat(difference).mean)/3
        assert error < 18, (path.name, 'photo still present in background', box, error)
    if needs_text:
        text = all_text(path)
        assert '轻转图片独立编辑' in text, (path.name, text)
        assert 'Editable photo objects' in text, (path.name, text)
    print('PASS', path.name, len(matches), 'separate photo media; repaired background; valid OOXML', flush=True)
    return {'file': path.name, 'photo_count': len(matches), 'picture_count': len(objects)}


def move_replace(path, output):
    """Exercise the same independent object edits that the user needs in PPT."""
    presentation = Presentation(path)
    slide = presentation.slides[0]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                and shape.width*shape.height < presentation.slide_width*presentation.slide_height*.8]
    assert len(pictures) >= 3
    first, second = pictures[:2]
    first.left += int(presentation.slide_width*.055)
    first.top += int(presentation.slide_height*.06)
    original = (second.left, second.top, second.width, second.height)
    replacement = io.BytesIO()
    Image.new('RGB', (270, 285), (32, 170, 110)).save(replacement, 'PNG')
    second._element.getparent().remove(second._element)
    replacement.seek(0)
    added = slide.shapes.add_picture(replacement, *original)
    added.name = 'Replacement picture acceptance'
    # Cropping is an independent picture property, unrelated to other photos.
    added.crop_left, added.crop_top = .1, .05
    presentation.save(output)
    reopened = Presentation(output)
    assert any(shape.name == added.name and abs(shape.crop_left-.1)<.0001
               for shape in reopened.slides[0].shapes)
    print('PASS move, replace and crop independently:', output.name, flush=True)


def native_regression(folder):
    pdf = folder/'native.pdf'
    transparent = Image.new('RGBA', (160, 160), (0, 0, 0, 0))
    ImageDraw.Draw(transparent).ellipse((8, 8, 152, 152), fill=(210, 65, 70, 190))
    stream = io.BytesIO()
    transparent.save(stream, 'PNG')
    with fitz.open() as document:
        page = document.new_page(width=600, height=400)
        page.insert_text((35, 60), 'Native PDF text remains editable', fontsize=22)
        page.insert_image(fitz.Rect(70, 100, 230, 260), stream=stream.getvalue())
        page.insert_image(fitz.Rect(145, 150, 415, 330), stream=stream.getvalue())
        document.save(pdf)
    for fmt in ('pptx', 'docx'):
        output = convert(pdf, fmt, mode='layout')[0]
        assert 'Native PDF text remains editable' in all_text(output)
        images, _ = picture_objects(output)
        assert len(images) >= 2, (fmt, 'overlapping native images collapsed')
        assert any(Image.open(io.BytesIO(obj['data'])).mode == 'RGBA' for obj in images), (
            fmt, 'native image alpha lost')
        print('PASS native text, overlapping images and alpha:', output.name, flush=True)


def full_photo_regression(folder):
    """A page that is one photograph needs no text to be a valid editable image."""
    image_pdf(folder/'full-photo.png', folder/'full-photo-scan.pdf')
    for fmt in ('pptx', 'docx'):
        output = convert(folder/'full-photo-scan.pdf', fmt, mode='ocr')[0]
        if fmt == 'pptx':
            deck = Presentation(output)
            assert len(deck.slides) == 1
            pictures = [shape for shape in deck.slides[0].shapes
                        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            originals = [shape for shape in pictures if shape.name == '可替换图片']
            assert len(originals) == 1, 'full photograph needs one intact independently editable asset'
            assert not all_text(output).strip(), 'text-free photograph acquired hallucinated OCR text'
            originals[0].crop_left = .08
            deck.save(folder/'full-photo-cropped.pptx')
        else:
            with zipfile.ZipFile(output) as archive:
                doc = ET.fromstring(archive.read('word/document.xml'))
                assert doc.findall('.//a:blip', NS), 'full photograph was discarded without OCR text'
                originals = [anchor for anchor in doc.findall('.//wp:anchor', NS)
                             if anchor.find('wp:docPr', NS).get('name') == '可替换图片']
                assert len(originals) == 1, 'full photograph needs one intact independently editable asset'
                assert not all_text(output).strip(), 'text-free photograph acquired hallucinated OCR text'
        print('PASS force OCR preserves photograph without editable text:', output.name, flush=True)


def native_collage_regression(folder):
    """A native photo-only PDF already has separate assets and must keep them."""
    source = folder/'native-photo-collage.pdf'
    with fitz.open() as document:
        page = document.new_page(width=900, height=600)
        for seed, x in enumerate((35, 330, 625), 1):
            stream = io.BytesIO()
            texture(seed).save(stream, 'PNG')
            page.insert_image(fitz.Rect(x, 110, x+230, 360), stream=stream.getvalue())
        document.save(source)
    for fmt in ('pptx', 'docx'):
        output = convert(source, fmt, mode='layout')[0]
        objects, _ = picture_objects(output)
        assert len(objects) == 3, (fmt, 'native photos needlessly merged/re-OCRed', len(objects))
        assert len({obj['data'] for obj in objects}) == 3
        print('PASS photo-only native collage retains three original assets:', output.name, flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--runtime', type=Path)
    parser.add_argument('--output-dir', type=Path, help='Keep fixtures and exports for visual verification')
    args = parser.parse_args()
    if args.runtime:
        import file_tools
        file_tools.tool_runtime = lambda: args.runtime.resolve()
    temp = None
    if args.output_dir:
        folder = args.output_dir.resolve()
        folder.mkdir(parents=True, exist_ok=True)
    else:
        temp = tempfile.TemporaryDirectory(prefix='.image-object-test-', dir=ROOT)
        folder = Path(temp.name)
    try:
        boxes, size, base = fixture(folder)
        summary = []
        for source_name in ('mixed-image.png', 'mixed-scan.pdf'):
            for fmt in ('pptx', 'docx'):
                print('Converting', source_name, '->', fmt, flush=True)
                path = convert(folder/source_name, fmt, mode='layout')[0]
                summary.append(acceptance(path, boxes, size, base))
                if fmt == 'pptx' and source_name == 'mixed-scan.pdf':
                    move_replace(path, folder/'mixed-scan-moved-replaced.pptx')
        for source_name, fmt in (('photo-only.png', 'pptx'), ('photo-only-scan.pdf', 'docx')):
            print('Converting text-free photos', source_name, '->', fmt, flush=True)
            path = convert(folder/source_name, fmt, mode='layout')[0]
            summary.append(acceptance(path, boxes, size, base, needs_text=False))
        native_regression(folder)
        native_collage_regression(folder)
        full_photo_regression(folder)
        (folder/'image-acceptance.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
        print('PASS: image/PDF -> PPT/Word; text-free photos; independent move/replace/crop; native PDF', flush=True)
    finally:
        if temp:
            temp.cleanup()


if __name__ == '__main__':
    main()
