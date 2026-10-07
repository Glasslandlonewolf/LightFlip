"""Paragraph boundaries and genuine editable Office structures, without OCR."""
from pathlib import Path
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from paragraph_layout import group_paragraphs, add_textbox
from pdf_layout import anchored_text
from pptx import Presentation
from pptx.util import Pt
from docx import Document


def line(text, x, y, width=220, **extra):
    return dict(kind='text', text=text, bbox=(x,y,x+width,y+18),
                origin=(x,y+20), size=22, font='Arial', flags=0,
                color=0x333333, angle=0, seq=y*1000+x, **extra)


def paragraphs(rows):
    return [o for o in group_paragraphs(rows,960,540) if o['kind']=='text']


def main():
    # Reading order interleaves two columns; each column has two paragraphs.
    rows=[line('LEFT first line of paragraph',40,60),line('LEFT continued text here',40,88),
          line('LEFT second paragraph begins',40,160),line('LEFT second continued',40,188),
          line('RIGHT first line of paragraph',500,60),line('RIGHT continued text here',500,88),
          line('RIGHT second paragraph begins',500,160),line('RIGHT second continued',500,188)]
    groups=paragraphs(rows)
    assert len(groups)==4 and all(len(o['paragraph_rows'])==2 for o in groups)
    assert all(not ('LEFT' in o['text'] and 'RIGHT' in o['text']) for o in groups)
    # Additional paragraph spacing across native blocks remains a boundary.
    native=[line('First native paragraph',40,60,block_id=1,line_id=0),
            line('Second native paragraph',40,95,block_id=2,line_id=0)]
    assert len(paragraphs(native))==2
    native[1]['bbox']=(40,86,260,104);native[1]['origin']=(40,106)
    assert len(paragraphs(native))==1  # PDF exporters may split every visual line.
    # Short final lowercase word uses the same font line, despite shorter ink.
    short=line('urna.',40,116,width=50);short['bbox']=(40,116,90,127)
    assert len(paragraphs(rows[:2]+[short]))==1
    # Headings, bullet items, panel changes, and graphical dividers stay separate.
    heading=line('Heading',40,32);heading['flags']=16
    assert len(paragraphs([heading]+rows[:2]))==2
    assert len(paragraphs([line('• First separate item',40,60),line('• Next separate item',40,88)]))==2
    assert len(paragraphs([line('First card body line',40,60,panel_color=0xFFFFFF),
                           line('Second card body line',40,88,panel_color=0x663333)]))==2
    divider=dict(kind='shape',bbox=(40,84,260,84),geometry='line',seq=0)
    assert len(paragraphs(rows[:2]+[divider]))==2
    panel=dict(kind='shape',bbox=(20,40,300,140),geometry='rect',seq=0)
    assert len(paragraphs(rows[:2]+[panel]))==1
    # Native bold spans remain runs inside one paragraph, not separate boxes.
    first=line('Native regular ',40,60,width=145,block_id=5,line_id=0)
    bold=line('emphasis',185,60,width=100,block_id=5,line_id=0);bold['flags']=16
    continued=line('Native continuation here',40,88,block_id=5,line_id=1)
    mixed=paragraphs([first,bold,continued])[0]
    assert len(mixed['paragraph_rows'])==2 and len(mixed['paragraph_rows'][0]['runs'])==2
    # An occasional OCR weight guess must not turn an entire body bold.
    noisy=[line('Paragraph ordinary opening',40,60,width_ratio=1,font_fit=.9),
           line('Paragraph noisy middle',40,88,width_ratio=1,font_fit=.9),
           line('Paragraph ordinary ending',40,116,width_ratio=1,font_fit=.9)]
    noisy[1]['flags']=16
    stable=paragraphs(noisy)[0]
    assert all(r['flags']==0 for r in stable['paragraph_rows'])
    with tempfile.TemporaryDirectory(prefix='.paragraph-test-', dir=Path.cwd()) as folder:
        folder=Path(folder)
        presentation=Presentation();presentation.slide_width=Pt(960);presentation.slide_height=Pt(540)
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        shapes=[add_textbox(slide,o) for o in groups+[mixed]]
        assert len(shapes)==5
        for shape in shapes:
            assert len(shape.text_frame.paragraphs)==1 and shape.text_frame.word_wrap
        # Edit and move the full paragraph as one object; its neighbour is intact.
        neighbour=shapes[1].text
        shapes[0].text_frame.paragraphs[0].runs[0].text+=' appended words for wrapping'
        shapes[0].left+=Pt(8)
        assert shapes[1].text==neighbour
        ppt=folder/'paragraphs.pptx';presentation.save(ppt)
        document=Document();p=document.add_paragraph()
        for index,obj in enumerate(groups+[mixed]):anchored_text(p,obj,1,index+1)
        doc=folder/'paragraphs.docx';document.save(doc)
        A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
        W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        with zipfile.ZipFile(ppt) as package:
            xml=ET.fromstring(package.read('ppt/slides/slide1.xml'))
            assert len(xml.findall('.//'+A+'p'))==5
            assert len(xml.findall('.//'+A+'br'))==5
        with zipfile.ZipFile(doc) as package:
            xml=ET.fromstring(package.read('word/document.xml'))
            contents=xml.findall('.//'+W+'txbxContent')
            assert len(contents)==5 and all(len(c.findall(W+'p'))==1 for c in contents)
            assert len(xml.findall('.//'+W+'br'))==5
    print('PASS: columns, paragraphs, headings, panels, lists, short endings, inline emphasis, stable OCR style; one editable PPT/Word paragraph each')


if __name__=='__main__':main()
