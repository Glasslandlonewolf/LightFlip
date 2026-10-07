"""Fast pixel acceptance for a non-text symbol inside a scanned circle.

Run with the prepared main Python environment. --runtime locates the separate
OCR Python environment; only artwork separation runs there, without OCR model
loading. No external image, user account, fixed workspace path, or Office app is
required. --output-dir keeps the generated files for optional visual rendering.
"""
from pathlib import Path
import argparse
import io
import json
import os
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from PIL import Image, ImageDraw, ImageChops, ImageStat


def fixture(folder):
    source = Image.new('RGB', (700, 500), (52, 63, 78))
    draw = ImageDraw.Draw(source)
    draw.ellipse((280, 180, 420, 320), fill=(251, 247, 232))
    draw.polygon([(350, 207), (382, 236), (363, 236), (363, 264), (378, 264),
                  (350, 294), (322, 264), (337, 264), (337, 236), (318, 236)],
                 fill=(114, 41, 53))
    source.save(folder/'circle-symbol.png')


def prepare(folder):
    import cv2
    import numpy as np
    from image_layout import separate_artwork
    source = folder/'circle-symbol.png'
    image = cv2.imdecode(np.fromfile(str(source), dtype='uint8'), cv2.IMREAD_COLOR)
    background = folder/'clean-circle.png'
    ok, encoded = cv2.imencode('.png', image)
    assert ok
    encoded.tofile(str(background))
    lines, images, shapes = separate_artwork(image, [], background, folder)
    item = {'width': 700, 'height': 500, 'page_width': 525., 'page_height': 375.,
            'lines': lines, 'images': images, 'shapes': shapes,
            'background': str(background), 'text': ''}
    (folder/'layout.json').write_text(json.dumps(item), encoding='utf-8')


def symbol_pixels(image):
    image = image.convert('RGBA')
    pixels = image.get_flattened_data() if hasattr(image, 'get_flattened_data') else image.getdata()
    return sum(abs(r-114)<8 and abs(g-41)<8 and abs(b-53)<8 and a>240 for r,g,b,a in pixels)


def export_and_verify(folder):
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from pdf_ocr import image_pdf
    from pdf_layout import layout_slides, layout_word
    from smoke_image_objects import picture_objects
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    item = json.loads((folder/'layout.json').read_text(encoding='utf-8'))
    image_pdf(folder/'circle-symbol.png', folder/'circle-symbol.pdf')
    expected = symbol_pixels(Image.open(folder/'circle-symbol.png'))
    report = []
    for fmt, writer in [('pptx', layout_slides), ('docx', layout_word)]:
        path = folder/('circle-symbol.'+fmt)
        writer(folder/'circle-symbol.pdf', path, ocr_pages={0: item})
        objects, background = picture_objects(path)
        retained = sum(symbol_pixels(Image.open(io.BytesIO(obj['data']))) for obj in objects)
        assert retained >= expected*.92, (fmt, 'symbol was erased by an empty ellipse', retained, expected)
        cleaned = Image.open(io.BytesIO(background['data'])).convert('RGB').resize((700, 500))
        blue = Image.new('RGB', (700, 500), (52, 63, 78))
        inner = (310, 210, 390, 290)
        error = sum(ImageStat.Stat(ImageChops.difference(cleaned.crop(inner), blue.crop(inner))).mean)/3
        assert error < 8, (fmt, 'duplicate symbol remains in background', error)
        report.append({'format': fmt, 'retained_symbol_pixels': retained,
                       'expected_symbol_pixels': expected, 'background_error': error})
        print('PASS non-text circle artwork:', fmt, 'retained symbol pixels', retained,
              'expected', expected, 'background error', error, flush=True)
    deck = Presentation(folder/'circle-symbol.pptx')
    artwork = [shape for shape in deck.slides[0].shapes
               if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and '背景' not in shape.name]
    assert artwork, 'symbol needs an independently movable picture'
    artwork[0].left += int(deck.slide_width*.15)
    deck.save(folder/'circle-symbol-moved.pptx')
    (folder/'acceptance.json').write_text(json.dumps(report, indent=2), encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--runtime', type=Path, help='Directory containing tools-venv or ocr-runtime')
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--prepare', type=Path, help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.prepare:
        prepare(args.prepare.resolve())
        return
    from file_tools import tool_runtime
    runtime = args.runtime.resolve() if args.runtime else tool_runtime()
    interpreter = runtime/'tools-venv/Scripts/python.exe'
    if not interpreter.is_file():
        interpreter = runtime/'ocr-runtime/python.exe'
    assert interpreter.is_file(), 'prepared OCR Python runtime is required'
    temporary = None
    if args.output_dir:
        folder = args.output_dir.resolve()
        folder.mkdir(parents=True, exist_ok=True)
    else:
        temporary = tempfile.TemporaryDirectory(prefix='.circle-artwork-test-', dir=ROOT)
        folder = Path(temporary.name)
    try:
        fixture(folder)
        env = os.environ.copy()
        env.pop('PYTHONHOME', None)
        env.pop('PYTHONPATH', None)
        env['PYTHONIOENCODING'] = 'utf-8'
        subprocess.run([str(interpreter), '-I', '-X', 'utf8', str(Path(__file__).resolve()),
                        '--prepare', str(folder)], check=True, env=env,
                       creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        export_and_verify(folder)
    finally:
        if temporary:
            temporary.cleanup()


if __name__ == '__main__':
    main()
