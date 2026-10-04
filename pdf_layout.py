"""Rebuild native PDF text and image objects at their original page positions."""
import io
import re
import math

def image_block_png(block, rect):
    """Render only the original image, never a composite crop of the page."""
    from PIL import Image, ImageChops
    image = Image.open(io.BytesIO(block['image'])).convert('RGBA')
    if block.get('mask'):
        mask = Image.open(io.BytesIO(block['mask'])).convert('L')
        if mask.size != image.size:
            mask = mask.resize(image.size)
        image.putalpha(ImageChops.multiply(image.getchannel('A'), mask))
    a,b,c,d,e,f = block['transform']
    determinant = a*d-b*c
    if abs(determinant) < 1e-10:
        raise ValueError('PDF 图片变换无效，无法安全拆分。')
    scale = min(2, 3000/max(rect.width,rect.height))
    width,height = max(1,math.ceil(rect.width*scale)),max(1,math.ceil(rect.height*scale))
    sx,sy = width/rect.width,height/rect.height
    iw,ih = image.size
    matrix = (iw*d/determinant/sx, -iw*c/determinant/sy,
              iw*(d*(rect.x0-e)-c*(rect.y0-f))/determinant,
              -ih*b/determinant/sx, ih*a/determinant/sy,
              ih*(-b*(rect.x0-e)+a*(rect.y0-f))/determinant)
    image = image.transform((width,height),Image.Transform.AFFINE,matrix,Image.Resampling.BICUBIC)
    stream=io.BytesIO()
    image.save(stream,'PNG')
    return stream.getvalue()

def drawing_png(drawing, width, height):
    """Keep vector overlays separate, with the PDF's original paint order."""
    import pymupdf as fitz
    with fitz.open() as document:
        page=document.new_page(width=width,height=height)
        shape=page.new_shape()
        for item in drawing['items']:
            kind=item[0]
            if kind=='l': shape.draw_line(item[1],item[2])
            elif kind=='c': shape.draw_bezier(item[1],item[2],item[3],item[4])
            elif kind=='re': shape.draw_rect(item[1])
            elif kind=='qu': shape.draw_quad(item[1])
            else: raise ValueError('PDF 含有暂不支持的矢量图形，请选择保留页面外观。')
        caps=drawing.get('lineCap') or (0,0,0)
        shape.finish(color=drawing.get('color'),fill=drawing.get('fill'),width=drawing.get('width') or 1,
                     lineCap=max(caps) if isinstance(caps,tuple) else caps,lineJoin=int(drawing.get('lineJoin') or 0),
                     dashes=drawing.get('dashes'),closePath=bool(drawing.get('closePath')),even_odd=bool(drawing.get('even_odd')),
                     fill_opacity=1 if drawing.get('fill_opacity') is None else drawing['fill_opacity'],
                     stroke_opacity=1 if drawing.get('stroke_opacity') is None else drawing['stroke_opacity'])
        shape.commit()
        rect=fitz.Rect(drawing['rect'])
        padding=max(1,(drawing.get('width') or 0)*5)
        rect=fitz.Rect(rect.x0-padding,rect.y0-padding,rect.x1+padding,rect.y1+padding)&page.rect
        if rect.is_empty:return None
        factor=min(2,3000/max(rect.width,rect.height))
        data=page.get_pixmap(matrix=fitz.Matrix(factor,factor),clip=rect,alpha=True,annots=False).tobytes('png')
        colors=[color for color in (drawing.get('color'),drawing.get('fill')) if color is not None]
        if colors and all(all(component>.999 for component in color) for color in colors):
            # Office resamples RGB even in fully transparent PNG pixels. Black
            # transparent padding otherwise creates gray halos around white shapes.
            from PIL import Image
            image=Image.open(io.BytesIO(data)).convert('RGBA')
            clean=Image.new('RGBA',image.size,'white')
            clean.putalpha(image.getchannel('A'))
            stream=io.BytesIO()
            clean.save(stream,'PNG')
            data=stream.getvalue()
        result={'kind':'image','bbox':tuple(rect),'data':data,'vector':True,'seq':drawing['seqno']}
        if len(drawing['items'])==1 and drawing['items'][0][0]=='re' and all((drawing.get(k) is None or drawing[k]==1) for k in ['fill_opacity','stroke_opacity']):
            result['native_rect']={'bbox':tuple(drawing['items'][0][1]),'color':drawing.get('color'),'fill':drawing.get('fill'),'width':drawing.get('width') or 1}
        return result

def font_name(name):
    name = re.sub(r'^[A-Z]{6}\+', '', name)
    for token, replacement in [('MicrosoftYaHei', 'Microsoft YaHei'), ('SimHei', '黑体'), ('SimSun', '宋体'), ('Arial', 'Arial'), ('Calibri', 'Calibri'), ('TimesNewRoman', 'Times New Roman')]:
        if token.lower() in name.lower():
            return replacement
    return re.sub(r'[-,](Bold|Italic|Regular).*$', '', name, flags=re.I)

def extract_pages(src):
    import pymupdf as fitz
    with fitz.open(src) as document:
        if document.needs_pass:
            raise ValueError('PDF 有密码保护，请先解锁。')
        for index in range(len(document)):
            # Work only in memory; normalize rotation without touching original PDF.
            with fitz.open() as copy:
                copy.insert_pdf(document, from_page=index, to_page=index)
                page = copy[0]
                page.remove_rotation()
                objects = []
                bboxlog=page.get_bboxlog()
                image_events=[(n,fitz.Rect(box)) for n,(kind,box) in enumerate(bboxlog) if kind=='fill-image']
                used_image_events=set()
                traces=page.get_texttrace()
                for block in page.get_text('rawdict')['blocks']:
                    if block['type'] == 1:
                        rect = fitz.Rect(block['bbox']) & page.rect
                        if rect.is_empty:
                            continue
                        available=[(n,r) for n,r in image_events if n not in used_image_events]
                        seq=min(available,key=lambda item:sum(abs(item[1][i]-block['bbox'][i]) for i in range(4)))[0] if available else block['number']
                        used_image_events.add(seq)
                        objects.append({'kind': 'image', 'bbox': tuple(rect), 'data': image_block_png(block,rect),'seq':seq})
                    elif block['type'] == 0:
                        for line in block['lines']:
                            for span in line['spans']:
                                chars = span['chars']
                                start, end = 0, len(chars)
                                while start < end and chars[start]['c'].isspace():
                                    start += 1
                                while end > start and chars[end-1]['c'].isspace():
                                    end -= 1
                                if start == end or span.get('alpha', 255) == 0:
                                    continue
                                kept = chars[start:end]
                                rect = fitz.Rect(kept[0]['bbox'])
                                for char in kept[1:]:
                                    rect |= fitz.Rect(char['bbox'])
                                text = ''
                                for position, char in enumerate(kept):
                                    if position and abs(line['dir'][1]) < .01:
                                        gap = char['bbox'][0]-kept[position-1]['bbox'][2]
                                        if gap > span['size']*.18 and not char['c'].isspace() and not kept[position-1]['c'].isspace():
                                            text += ' '*max(1, round(gap/(span['size']*.25)))
                                    text += char['c']
                                item = {k: span[k] for k in ['size', 'flags', 'color', 'font']}
                                item.update(kind='text', bbox=tuple(rect), text=text, origin=kept[0]['origin'], angle=math.degrees(math.atan2(line['dir'][1], line['dir'][0])))
                                candidates=[]
                                for trace in traces:
                                    for char in trace['chars']:
                                        if char[0]==ord(kept[0]['c']):
                                            distance=(char[2][0]-kept[0]['origin'][0])**2+(char[2][1]-kept[0]['origin'][1])**2
                                            candidates.append((distance,trace['seqno']))
                                item['seq']=min(candidates)[1] if candidates else len(bboxlog)
                                objects.append(item)
                for drawing in page.get_drawings():
                    obj=drawing_png(drawing,page.rect.width,page.rect.height)
                    if obj:objects.append(obj)
                # Extracted pictures are independent original image data. Never render
                # the page while temporary redaction annotations are present.
                page.add_redact_annot(page.rect, fill=False, cross_out=False)
                page.apply_redactions(images=1, graphics=2, text=0)
                background = page.get_pixmap(matrix=fitz.Matrix(min(2, 3000/max(page.rect.width, page.rect.height)), min(2, 3000/max(page.rect.width, page.rect.height))), alpha=False,annots=False).tobytes('png')
                objects.sort(key=lambda obj:obj['seq'])
                yield {'width': page.rect.width, 'height': page.rect.height, 'objects': objects, 'background': background}

def export_layout(src, dst, fmt):
    if fmt == 'pptx':
        layout_slides(src, dst)
    else:
        layout_word(src, dst)

def layout_slides(src, dst):
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    presentation = Presentation()
    for index, page in enumerate(extract_pages(src)):
        if index == 0:
            scale = min(1, 2800/max(page['width'], page['height']))
            presentation.slide_width = Pt(max(72, page['width']*scale))
            presentation.slide_height = Pt(max(72, page['height']*scale))
        scale = min(presentation.slide_width/12700/page['width'], presentation.slide_height/12700/page['height'])
        ox = (presentation.slide_width/12700-page['width']*scale)/2
        oy = (presentation.slide_height/12700-page['height']*scale)/2
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.shapes.add_picture(io.BytesIO(page['background']), Pt(ox), Pt(oy), width=Pt(page['width']*scale), height=Pt(page['height']*scale))
        background.line.fill.background()
        background.name = '页面图形背景（文字和图片已分离）'
        for obj in page['objects']:
            x0, y0, x1, y1 = obj['bbox']
            if obj['kind'] == 'image':
                if obj.get('native_rect'):
                    from pptx.enum.shapes import MSO_SHAPE
                    rect=obj['native_rect']
                    rx0,ry0,rx1,ry1=rect['bbox']
                    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(ox+rx0*scale),Pt(oy+ry0*scale),Pt((rx1-rx0)*scale),Pt((ry1-ry0)*scale))
                    if rect['fill'] is not None:
                        shape.fill.solid()
                        shape.fill.fore_color.rgb=RGBColor(*(round(c*255) for c in rect['fill']))
                    else:shape.fill.background()
                    if rect['color'] is not None:
                        shape.line.color.rgb=RGBColor(*(round(c*255) for c in rect['color']))
                        shape.line.width=Pt(rect['width']*scale)
                    else:shape.line.fill.background()
                    # PDF paths have no Office theme effects. A default rectangle
                    # inherits the presentation theme's shadow even with white fill.
                    from pptx.oxml.xmlchemy import OxmlElement
                    from pptx.oxml.ns import qn
                    style = shape._element.find(qn('p:style'))
                    if style is not None:
                        shape._element.remove(style)
                    shape._element.spPr.append(OxmlElement('a:effectLst'))
                else:
                    shape = slide.shapes.add_picture(io.BytesIO(obj['data']), Pt(ox+x0*scale), Pt(oy+y0*scale), width=Pt((x1-x0)*scale), height=Pt((y1-y0)*scale))
                    shape.line.fill.background()
                shape.name = '图形标注' if obj.get('vector') else '可替换图片'
            else:
                top = obj['origin'][1] - obj['size']*.92
                shape = slide.shapes.add_textbox(Pt(ox+x0*scale), Pt(oy+top*scale), Pt(max(x1-x0+3, obj['size'])*scale), Pt(max(y1-y0+3, obj['size']*1.6)*scale))
                shape.line.fill.background()
                shape.fill.background()
                shape.name = '可编辑文字'
                shape.rotation = obj['angle']
                frame = shape.text_frame
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                frame.word_wrap = False
                frame.auto_size = MSO_AUTO_SIZE.NONE
                frame.vertical_anchor = MSO_ANCHOR.TOP
                p = frame.paragraphs[0]
                p.space_before = p.space_after = Pt(0)
                run = p.add_run()
                run.text = obj['text']
                run.font.name = font_name(obj['font'])
                from pptx.oxml.xmlchemy import OxmlElement
                for name in ['a:ea', 'a:cs']:
                    element = OxmlElement(name)
                    element.set('typeface', font_name(obj['font']))
                    run._r.get_or_add_rPr().append(element)
                run.font.size = Pt(obj['size']*scale)
                run.font.bold = bool(obj['flags'] & 16)
                run.font.italic = bool(obj['flags'] & 2)
                c = obj['color']
                run.font.color.rgb = RGBColor((c>>16)&255, (c>>8)&255, c&255)
    presentation.save(dst)

def anchored_picture(paragraph, data, rect, behind=False, vector=False, layer=100):
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    x, y, width, height = rect
    inline = paragraph.add_run().add_picture(io.BytesIO(data), width=Pt(width), height=Pt(height))._inline
    anchor = OxmlElement('wp:anchor')
    for key, value in {'distT':'0', 'distB':'0', 'distL':'0', 'distR':'0', 'simplePos':'0', 'relativeHeight':'0' if behind else str(layer), 'behindDoc':'1' if behind else '0', 'locked':'0', 'layoutInCell':'1', 'allowOverlap':'1'}.items():
        anchor.set(key, value)
    simple = OxmlElement('wp:simplePos')
    simple.set('x', '0')
    simple.set('y', '0')
    anchor.append(simple)
    for axis, offset in [('H', x), ('V', y)]:
        position = OxmlElement('wp:position'+axis)
        position.set('relativeFrom', 'page')
        value = OxmlElement('wp:posOffset')
        value.text = str(round(offset*12700))
        position.append(value)
        anchor.append(position)
    anchor.append(inline.extent)
    anchor.append(OxmlElement('wp:wrapNone'))
    for child in list(inline):
        if child.tag.endswith('}docPr') or child.tag.endswith('}cNvGraphicFramePr') or child.tag.endswith('}graphic'):
            if child.tag.endswith('}docPr'):
                child.set('name', '页面图形背景' if behind else '图形标注' if vector else '可替换图片')
                child.set('descr', '图形背景，文字和独立图片已分离' if behind else '可以移动、缩放或替换的独立图片')
            anchor.append(child)
    inline.getparent().replace(inline, anchor)

def anchored_text(paragraph, obj, scale, identifier):
    from lxml import etree
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    V='urn:schemas-microsoft-com:vml'
    W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    pict = OxmlElement('w:pict')
    shape = etree.SubElement(pict, '{'+V+'}shape', nsmap={'v':V})
    shape.set('id', 'LightFlipText'+str(identifier))
    shape.set('type', '#_x0000_t202')
    shape.set('filled', 'f')
    shape.set('stroked', 'f')
    x0,y0,x1,y1=obj['bbox']
    top=obj['origin'][1]-obj['size']*.92
    shape.set('style', f'position:absolute;margin-left:{x0*scale:.3f}pt;margin-top:{top*scale:.3f}pt;width:{max(x1-x0+3,obj["size"])*scale:.3f}pt;height:{max(y1-y0+3,obj["size"]*1.8)*scale:.3f}pt;rotation:{obj["angle"]:.3f};z-index:{obj.get("seq",100)+100};mso-position-horizontal-relative:page;mso-position-vertical-relative:page;mso-wrap-style:none')
    textbox=etree.SubElement(shape,'{'+V+'}textbox')
    textbox.set('inset', '0,0,0,0')
    content=etree.SubElement(textbox,'{'+W+'}txbxContent')
    p=etree.SubElement(content,'{'+W+'}p')
    props=etree.SubElement(p,'{'+W+'}pPr')
    for key in ['autoSpaceDE','autoSpaceDN']:
        element=etree.SubElement(props,'{'+W+'}'+key)
        element.set(qn('w:val'),'0')
    spacing=etree.SubElement(props,'{'+W+'}spacing')
    spacing.set(qn('w:before'),'0')
    spacing.set(qn('w:after'),'0')
    spacing.set(qn('w:line'),str(round(obj['size']*scale*1.2*20)))
    spacing.set(qn('w:lineRule'),'exact')
    run=etree.SubElement(p,'{'+W+'}r')
    rpr=etree.SubElement(run,'{'+W+'}rPr')
    fonts=etree.SubElement(rpr,'{'+W+'}rFonts')
    for key in ['ascii','hAnsi','eastAsia','cs']:
        fonts.set(qn('w:'+key),font_name(obj['font']))
    sz=etree.SubElement(rpr,'{'+W+'}sz')
    sz.set(qn('w:val'),str(max(2,round(obj['size']*scale*2))))
    if obj['flags'] & 16:
        etree.SubElement(rpr,'{'+W+'}b')
    if obj['flags'] & 2:
        etree.SubElement(rpr,'{'+W+'}i')
    color=etree.SubElement(rpr,'{'+W+'}color')
    color.set(qn('w:val'),f'{obj["color"]:06X}')
    text=etree.SubElement(run,'{'+W+'}t')
    text.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    text.text=obj['text']
    paragraph.add_run()._r.append(pict)


def anchored_rectangle(paragraph, obj, scale, identifier):
    """Use the same DrawingML paint order as pictures, without PNG edge halos."""
    from lxml import etree
    from docx.oxml.ns import qn
    from PIL import Image
    WPS='http://schemas.microsoft.com/office/word/2010/wordprocessingShape'
    rect=obj['native_rect']
    x0,y0,x1,y1=rect['bbox']
    stream=io.BytesIO()
    Image.new('RGB',(1,1),'white').save(stream,'PNG')
    anchored_picture(paragraph,stream.getvalue(),(x0*scale,y0*scale,(x1-x0)*scale,(y1-y0)*scale),vector=True,layer=obj['seq']+100)
    anchor=paragraph.runs[-1]._r.find('.//'+qn('wp:anchor'))
    anchor.find(qn('wp:docPr')).set('name','页面矩形图形')
    graphic_data=anchor.find('.//'+qn('a:graphicData'))
    graphic_data.clear()
    graphic_data.set('uri',WPS)
    shape=etree.SubElement(graphic_data,'{'+WPS+'}wsp',nsmap={'wps':WPS})
    etree.SubElement(shape,'{'+WPS+'}cNvSpPr')
    props=etree.SubElement(shape,'{'+WPS+'}spPr')
    transform=etree.SubElement(props,qn('a:xfrm'))
    etree.SubElement(transform,qn('a:off'),x='0',y='0')
    extent=anchor.find(qn('wp:extent'))
    etree.SubElement(transform,qn('a:ext'),cx=extent.get('cx'),cy=extent.get('cy'))
    geometry=etree.SubElement(props,qn('a:prstGeom'),prst='rect')
    etree.SubElement(geometry,qn('a:avLst'))
    def color(value):
        return ''.join(f'{max(0,min(255,round(component*255))):02X}' for component in value)
    def fill(parent,value):
        if value is None:
            etree.SubElement(parent,qn('a:noFill'))
        else:
            solid=etree.SubElement(parent,qn('a:solidFill'))
            etree.SubElement(solid,qn('a:srgbClr'),val=color(value))
    fill(props,rect['fill'])
    line=etree.SubElement(props,qn('a:ln'),w=str(round(rect['width']*scale*12700)))
    fill(line,rect['color'])
    etree.SubElement(shape,'{'+WPS+'}bodyPr')

def layout_word(src,dst):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    document=Document()
    identifier=0
    for index,page in enumerate(extract_pages(src)):
        section=document.sections[0] if index==0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        scale=min(1,1500/max(page['width'],page['height']))
        section.page_width=Pt(page['width']*scale)
        section.page_height=Pt(page['height']*scale)
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Pt(0)
        section.header_distance=section.footer_distance=Pt(0)
        paragraph=document.add_paragraph()
        paragraph.paragraph_format.space_before=paragraph.paragraph_format.space_after=Pt(0)
        paragraph.paragraph_format.line_spacing=Pt(1)
        anchored_picture(paragraph,page['background'],(0,0,page['width']*scale,page['height']*scale),behind=True)
        for obj in page['objects']:
            if obj['kind']=='image':
                if obj.get('native_rect'):
                    identifier+=1
                    anchored_rectangle(paragraph,obj,scale,identifier)
                else:
                    x0,y0,x1,y1=obj['bbox']
                    anchored_picture(paragraph,obj['data'],(x0*scale,y0*scale,(x1-x0)*scale,(y1-y0)*scale),vector=obj.get('vector',False),layer=obj['seq']+100)
            else:
                identifier+=1
                anchored_text(paragraph,obj,scale,identifier)
    document.save(dst)

