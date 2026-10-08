"""Native packaged-app smoke test; synthetic fixtures only, no user material."""
import hashlib
import json
import platform
import subprocess
import tempfile
from pathlib import Path


def main(report):
    from app import App
    from core import convert,office_pdf,publish
    from file_tools import ai_job
    from platform_support import libreoffice
    from PIL import Image,ImageDraw,ImageFont
    from pptx import Presentation
    import zipfile
    from lxml import etree
    import fitz
    checks=[]
    gui=App();gui.root.withdraw();gui.root.update();gui.root.destroy()
    checks.append('GUI and native drag/drop library loaded')
    with tempfile.TemporaryDirectory(prefix='lightflip-self-test-') as temp:
        folder=Path(temp)
        pdf=folder/'sample.pdf'
        page=fitz.open();p=page.new_page(width=720,height=480)
        p.insert_text((40,80),'LightFlip Mac test',fontsize=26)
        for index,text in enumerate(('This paragraph stays together when edited.',
                                     'Each line belongs to the same text box.',
                                     'The original document must stay unchanged.')):
            p.insert_text((40,150+index*24),text,fontsize=16)
        image=Image.new('RGB',(160,160),'#425bdb')
        import io
        stream=io.BytesIO();image.save(stream,format='PNG')
        p.insert_image(fitz.Rect(520,180,660,320),stream=stream.getvalue())
        page.save(pdf);page.close()
        before=hashlib.sha256(pdf.read_bytes()).hexdigest()
        ppt=convert(pdf,'pptx',mode='layout')[0]
        word=convert(pdf,'docx',mode='layout')[0]
        slides=Presentation(ppt)
        body=[shape for shape in slides.slides[0].shapes if shape.has_text_frame and 'Each line' in shape.text]
        assert len(body)==1 and 'This paragraph' in body[0].text and 'original document' in body[0].text
        assert any(shape.shape_type==13 for shape in slides.slides[0].shapes)
        with zipfile.ZipFile(word) as package:
            xml=etree.fromstring(package.read('word/document.xml'))
            ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            paragraphs=[''.join(node.itertext()) for node in xml.findall('.//w:txbxContent/w:p',ns)]
            assert any('Each line' in text and 'This paragraph' in text for text in paragraphs)
        checks.append('PDF to editable paragraph PPT/Word with independent artwork')
        assert hashlib.sha256(pdf.read_bytes()).hexdigest()==before
        old=folder/'reserved.txt';old.write_text('original')
        staged=folder/'staged.txt';staged.write_text('new')
        output=publish(staged,folder,'reserved','txt')
        assert output!=old and old.read_text()=='original' and output.read_text()=='new'
        checks.append('Original files and existing output preserved')
        fontpath=next((p for p in (Path('/System/Library/Fonts/Supplemental/Arial.ttf'),
                                    Path('/Library/Fonts/Arial.ttf')) if p.is_file()),None)
        font=ImageFont.truetype(str(fontpath),40) if fontpath else ImageFont.load_default(size=40)
        scanned=Image.new('RGB',(1100,180),'white')
        ImageDraw.Draw(scanned).text((30,50),'LightFlip editable paragraph test',font=font,fill='black')
        source=folder/'ocr.png';scanned.save(source)
        import rapidocr
        assert len(list((Path(rapidocr.__file__).parent/'models').rglob('*.onnx')))>=3, 'Offline OCR models are missing'
        result=ai_job({'kind':'ocr_text','files':[str(source)]},folder,lambda text:None)
        assert 'lightflip' in result['items'][0]['text'].lower(),result
        checks.append('Bundled offline OCR via separate worker process')
        for fmt in ('pptx','docx'):
            editable=convert(source,fmt,mode='ocr')[0]
            if fmt=='pptx':
                contents=' '.join(s.text for s in Presentation(editable).slides[0].shapes if s.has_text_frame)
            else:
                with zipfile.ZipFile(editable) as package:
                    node=etree.fromstring(package.read('word/document.xml'))
                    contents=' '.join(node.xpath('//w:t/text()',namespaces=ns))
            assert 'lightflip' in contents.lower(),contents
        checks.append('Scanned image to editable PPT/Word using installed Mac fonts')
        from imageio_ffmpeg import get_ffmpeg_exe
        subprocess.run([get_ffmpeg_exe(),'-version'],check=True,capture_output=True)
        checks.append('Bundled FFmpeg started')
        if libreoffice():
            for source in (ppt,word):
                rendered=folder/(source.suffix[1:]+'.pdf')
                office_pdf(source,rendered,source.suffix[1:])
                with fitz.open(rendered) as pages:
                    assert len(pages)==1
                    assert 'paragraph' in pages[0].get_text().lower()
                    pages[0].get_pixmap().save(str(folder/(source.suffix[1:]+'.png')))
            checks.append('Native LibreOffice PPT/Word rendering')
    report.parent.mkdir(parents=True,exist_ok=True)
    report.write_text(json.dumps({'passed':True,'architecture':platform.machine(),'checks':checks},indent=2),encoding='utf-8')
