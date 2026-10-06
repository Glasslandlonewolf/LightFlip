"""Office output with automatic local OCR reconstruction for image PDFs."""
import io
import re
import unicodedata
from pathlib import Path

def page_images(src):
    if src.suffix.lower() == '.pdf':
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(str(src)) as document:
            if not len(document):
                raise ValueError('PDF 没有页面。')
            for index in range(len(document)):
                page = document[index]
                width, height = page.get_size()
                bitmap = page.render(scale=min(2.0, 3000/max(width, height)))
                try:
                    image = bitmap.to_pil()
                    stream = io.BytesIO()
                    image.save(stream, 'PNG')
                    stream.seek(0)
                    yield stream, width/72, height/72
                finally:
                    bitmap.close()
                    page.close()
    else:
        from PIL import Image, ImageOps
        with Image.open(src) as original:
            image = ImageOps.exif_transpose(original).convert('RGBA')
            stream = io.BytesIO()
            image.save(stream, 'PNG')
            stream.seek(0)
            yield stream, image.width/96, image.height/96

def extract_text(src,progress=None):
    if src.suffix.lower() in {'.txt', '.md'}:
        text = src.read_text(encoding='utf-8-sig')
    elif src.suffix.lower()=='.pdf' or src.suffix.lower()[1:] in {'png','jpg','jpeg','gif','webp','bmp','tif','tiff'}:
        from pdf_ocr import extract_document_text
        text=extract_document_text(src,progress)
    else:
        from markitdown import MarkItDown
        text = MarkItDown().convert_local(str(src)).text_content
        meaningful=re.sub(r'<!--.*?-->|!\[[^\]]*\]\([^)]*\)','',text,flags=re.S).strip()
        if not meaningful and src.suffix.lower()[1:] in {'docx','pptx','xlsx','xls'}:
            from core import office_available,office_pdf
            if office_available(src.suffix.lower()[1:]):
                import tempfile
                from pdf_ocr import extract_document_text
                with tempfile.TemporaryDirectory(prefix='.lightflip-office-text-',dir=src.parent) as folder:
                    pdf=Path(folder)/'pages.pdf'
                    office_pdf(src,pdf,src.suffix.lower()[1:])
                    text=extract_document_text(pdf,progress)
    if not text.strip():
        raise ValueError('没有可提取的原生文字。扫描件请使用“文件工具 → OCR 提取文字”；图纸可选择“保留页面外观”。')
    # Office XML excludes control characters, preserve tabs/newlines.
    return ''.join(c for c in text if c in '\n\r\t' or ord(c) >= 32)

def export_office(src, dst, fmt, mode, progress=None):
    if mode not in {'pages', 'text', 'layout', 'ocr'}:
        raise ValueError('未知转换模式。')
    if (src.suffix.lower()=='.pdf' or src.suffix.lower()[1:] in {'png','jpg','jpeg','gif','webp','bmp','tif','tiff'}) and mode in {'layout','ocr'}:
        from pdf_ocr import layout_with_ocr
        layout_with_ocr(src,dst,fmt,progress=progress,force=mode=='ocr')
        return
    if mode=='ocr':
        raise ValueError('图片文字重建适用于图片或 PDF 转 Word / PPT。')
    if mode == 'layout' and src.suffix.lower() == '.pdf':
        from pdf_layout import export_layout
        export_layout(src, dst, fmt)
        return
    visual = mode=='pages' and src.suffix.lower() in {'.pdf','.png','.jpg','.jpeg','.gif','.webp','.bmp','.tif','.tiff'}
    if fmt == 'docx':
        if visual:
            images_to_word(src, dst)
        else:
            text_to_word(extract_text(src,progress), dst)
    else:
        if visual:
            images_to_slides(src, dst)
        else:
            text_to_slides(extract_text(src,progress), src.stem, dst)

def images_to_word(src, dst):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.section import WD_SECTION_START
    document = Document()
    for index, (stream, width, height) in enumerate(page_images(src)):
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        # Word page maximum is 22 inches; keep original proportions.
        scale = min(1, 21/max(width, height))
        section.page_width = Inches(width*scale + .5)
        section.page_height = Inches(height*scale + .5)
        section.left_margin = section.right_margin = Inches(.25)
        section.top_margin = section.bottom_margin = Inches(.25)
        section.header_distance = section.footer_distance = Inches(0)
        if index:
            break_paragraph = document.paragraphs[-1]
            break_paragraph.paragraph_format.space_before = Pt(0)
            break_paragraph.paragraph_format.space_after = Pt(0)
            break_paragraph.paragraph_format.line_spacing = Pt(1)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        run = paragraph.add_run()
        run.font.size = Pt(1)
        factor = min((width*scale)/(width), (height*scale-.15)/height)
        run.add_picture(stream, width=Inches(width*factor), height=Inches(height*factor))
    document.save(dst)

def images_to_slides(src, dst):
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    for index, (stream, width, height) in enumerate(page_images(src)):
        if index == 0:
            scale = min(1, 40/max(width, height))
            presentation.slide_width = Inches(max(1, width*scale))
            presentation.slide_height = Inches(max(1, height*scale))
        canvas_w = presentation.slide_width/914400
        canvas_h = presentation.slide_height/914400
        factor = min(canvas_w/width, canvas_h/height)
        w, h = width*factor, height*factor
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(stream, Inches((canvas_w-w)/2), Inches((canvas_h-h)/2), width=Inches(w), height=Inches(h))
    presentation.save(dst)

def text_to_word(text, dst):
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = document.styles[style_name]
        style.font.name = 'Microsoft YaHei'
        style.font.size = Pt(11 if style_name == 'Normal' else 16)
        fonts = style.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            style.element.get_or_add_rPr().append(fonts)
        fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for line in text.splitlines():
        if line.startswith('<!--') and line.rstrip().endswith('-->'):
            continue
        match = re.match(r'^(#{1,6})\s+(.*)', line)
        if match:
            document.add_heading(match.group(2), min(len(match.group(1)), 3))
        else:
            document.add_paragraph(line)
    document.save(dst)

def wrapped_lines(text, limit=80):
    lines = []
    for raw in text.splitlines():
        if raw.startswith('<!--') and raw.rstrip().endswith('-->'):
            continue
        raw = re.sub(r'^#{1,6}\s+', '', raw).expandtabs(4)
        line, width = '', 0
        for char in raw:
            units = 2 if unicodedata.east_asian_width(char) in {'W', 'F'} else 1
            if width+units > limit:
                lines.append(line)
                line, width = '', 0
            line += char
            width += units
        lines.append(line)
    return lines

def text_to_slides(text, title, dst):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    lines = wrapped_lines(text)
    for start in range(0, len(lines), 13):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        heading = slide.shapes.add_textbox(Inches(.55), Inches(.35), Inches(12.2), Inches(.6)).text_frame
        heading.text = f'{title[:40]} · {start//13+1}'
        heading.paragraphs[0].font.size = Pt(24)
        heading.paragraphs[0].font.name = 'Microsoft YaHei'
        body = slide.shapes.add_textbox(Inches(.55), Inches(1.2), Inches(12.2), Inches(5.8)).text_frame
        body.word_wrap = False
        for i, line in enumerate(lines[start:start+13]):
            paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(20)
            paragraph.font.name = 'Microsoft YaHei'
            paragraph.font.color.rgb = RGBColor(37, 49, 77)
            paragraph.space_after = Pt(6)
    presentation.save(dst)
