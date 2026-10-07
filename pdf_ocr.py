"""Rebuild scanned pages as editable text, independent images and simple shapes."""
from pathlib import Path
import io
import math
import tempfile


def image_page(page, require_dominant=False):
    """Full image pages may still contain a native footer or watermark."""
    images=page.get_image_info()
    if not images:
        return False
    if not require_dominant and not page.get_text().strip():
        return True
    area=page.rect.width*page.rect.height
    return any(max(0,min(page.rect.x1,item['bbox'][2])-max(page.rect.x0,item['bbox'][0]))*
               max(0,min(page.rect.y1,item['bbox'][3])-max(page.rect.y0,item['bbox'][1]))/area>=.50
               for item in images)


def visible_native_content(page):
    """Preserve rich native text; a footer or hidden OCR layer is not enough."""
    baselines=[]
    boxes=[]
    characters=0
    for trace in page.get_texttrace():
        if trace.get('type',0)==3 or trace.get('opacity',1)<=0:
            continue
        tolerance=max(2,trace.get('size',8)*.25)
        for codepoint,_,origin,bbox in trace['chars']:
            if not 32<=codepoint<=0x10FFFF or chr(codepoint).isspace():
                continue
            characters+=1
            boxes.append(bbox)
            if not any(abs(origin[1]-baseline)<=tolerance for baseline in baselines):
                baselines.append(origin[1])
    if characters<80 or len(baselines)<3:
        return False
    height=max(box[3] for box in boxes)-min(box[1] for box in boxes)
    return height>=page.rect.height*.25


def prepare_pages(src, folder, progress, force=False):
    import pymupdf as fitz
    from file_tools import ai_job
    sources, numbers, sizes = [], [], []
    with fitz.open(src) as document:
        if document.needs_pass:
            raise ValueError('PDF 有密码保护，请先解锁。')
        for index, original in enumerate(document):
            with fitz.open() as normalized:
                normalized.insert_pdf(document,from_page=index,to_page=index)
                page=normalized[0]
                page.remove_rotation()
                # Small original PDF pictures are already independent objects.
                # Keep their native data instead of rasterising them for OCR.
                if not force:
                    if not image_page(page,require_dominant=True) or visible_native_content(page):
                        continue
                progress(f'正在准备图片页文字与图片分离 · 第 {index+1}/{len(document)} 页')
                factor=min(3,3200/max(page.rect.width,page.rect.height))
                image=folder/f'page-{index+1}.png'
                page.get_pixmap(matrix=fitz.Matrix(factor,factor),alpha=False,annots=False).save(image)
                sources.append(str(image))
                numbers.append(index)
                sizes.append((page.rect.width,page.rect.height))
    if not sources:
        return {}
    result=ai_job({'kind':'ocr_layout','files':sources,'output_dir':str(folder)},folder,progress)
    pages={}
    for index,size,item in zip(numbers,sizes,result['items']):
        if item.get('lines') or item.get('images') or item.get('shapes'):
            item['page_width'],item['page_height']=size
            pages[index]=item
        else:
            # A full-page photograph may contain no text or separable objects.
            # The native writer can still export that original picture as an
            # independently movable/croppable asset, so an empty OCR result is
            # not a conversion failure (also keep genuinely blank pages).
            progress(f'第 {index+1} 页未找到可重建的文字或图片区域，保留原页面对象。')
    return pages


def text_objects(page):
    """Translate OCR geometry into the native layout writer's object protocol."""
    from PIL import ImageFont
    from ocr_layout import fonts
    factor=page['page_width']/page['width']
    choices=fonts()
    objects=[]
    for index,row in enumerate(page.get('lines', [])):
        x0,y0,x1,y1=row['glyph_bbox']
        style=row['style']
        size=style['size_px']*factor
        match=next((path for family,path,bold,italic in choices if
                    (family,bold,italic)==(style['family'],style['bold'],style['italic'])),None)
        top=y0*factor
        if match:
            font=ImageFont.truetype(match,max(5,round(style['size_px'])))
            box=font.getbbox(row['text'])
            ascent,_=font.getmetrics()
            top-=(box[1]-ascent)*factor+size*.92
        red,green,blue=row['color']
        objects.append({'kind':'text','bbox':tuple(v*factor for v in (x0,y0,x1,y1)),
                        'origin':(x0*factor,top+size*.92),'size':size,
                        'font':style['family'],'flags':(16 if style['bold'] else 0)|(2 if style['italic'] else 0),
                        'color':(red<<16)|(green<<8)|blue,'text':row['text'],'seq':index,
                        'width_ratio':style.get('width_ratio',1),
                        'font_fit':style.get('fit',0),
                        'panel_color':sum(v<<shift for v,shift in zip(row.get('background_color',[255,255,255]),(16,8,0))),
                        'angle':math.degrees(math.atan2(row['points'][1][1]-row['points'][0][1],
                                                       row['points'][1][0]-row['points'][0][0]))})
    from paragraph_layout import group_paragraphs
    grouped=group_paragraphs(objects+graphic_objects(page),page['page_width'],page['page_height'])
    return [obj for obj in grouped if obj['kind']=='text']


def graphic_objects(page):
    """Use one object protocol for PPT and Word, with OCR pixel geometry scaled."""
    sx=page['page_width']/page['width']
    sy=page['page_height']/page['height']
    objects=[]
    for shape in page.get('shapes', []):
        x0,y0,x1,y1=shape['bbox']
        kind=shape['type']
        if kind not in {'rect', 'ellipse', 'line'}:
            continue
        color=shape.get('color', [0,0,0])
        fill=shape.get('fill', None if kind=='line' else color)
        stroke=shape.get('line_color', color if kind=='line' else None)
        rgb=lambda value:None if value is None else tuple(max(0,min(255,v))/255 for v in value)
        objects.append({'kind':'shape', 'bbox':(x0*sx,y0*sy,x1*sx,y1*sy),
                        'geometry':kind, 'fill':rgb(fill), 'color':rgb(stroke),
                        'width':shape.get('line_width',1)*sx,
                        'name':shape.get('name') or {'rect':'可编辑矩形','ellipse':'可编辑圆形','line':'可编辑线条'}[kind],
                        'seq':shape.get('seq',len(objects))})
    for image in page.get('images', []):
        x0,y0,x1,y1=image['bbox']
        objects.append({'kind':'image', 'bbox':(x0*sx,y0*sy,x1*sx,y1*sy),
                        'data':Path(image['path']).read_bytes(),
                        'name':image.get('name') or '独立图片（可移动、裁剪、替换）',
                        'seq':image.get('seq',len(objects))})
    objects.sort(key=lambda obj:obj['seq'])
    # A separate text layer stays above extracted graphics. This also avoids
    # Word's VML and DrawingML paint-order differences for overlapping anchors.
    for index,obj in enumerate(objects):
        obj['seq']=index
    return objects


def layout_objects(page):
    objects=graphic_objects(page)
    for obj in text_objects(page):
        obj['seq']=len(objects)
        objects.append(obj)
    return objects


def add_graphic(slide,obj,scale,ox,oy):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    x0,y0,x1,y1=obj['bbox']
    if obj['kind']=='image':
        shape=slide.shapes.add_picture(io.BytesIO(obj['data']),Pt(ox+x0*scale),Pt(oy+y0*scale),
                                      width=Pt((x1-x0)*scale),height=Pt((y1-y0)*scale))
        shape.line.fill.background()
    else:
        if obj['geometry']=='line':
            shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(ox+x0*scale),Pt(oy+y0*scale),
                                             Pt(ox+x1*scale),Pt(oy+y1*scale))
        else:
            geometry=MSO_SHAPE.OVAL if obj['geometry']=='ellipse' else MSO_SHAPE.RECTANGLE
            shape=slide.shapes.add_shape(geometry,Pt(ox+x0*scale),Pt(oy+y0*scale),
                                         Pt((x1-x0)*scale),Pt((y1-y0)*scale))
            if obj['fill'] is None:
                shape.fill.background()
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb=RGBColor(*(round(c*255) for c in obj['fill']))
        if obj['color'] is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb=RGBColor(*(round(c*255) for c in obj['color']))
            shape.line.width=Pt(obj['width']*scale)
        style=shape._element.find(qn('p:style'))
        if style is not None:
            shape._element.remove(style)
        shape._element.spPr.append(OxmlElement('a:effectLst'))
    shape.name=obj['name']
    return shape


def add_ocr_slide(presentation,page,scale,ox,oy):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR,MSO_AUTO_SIZE
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    image=slide.shapes.add_picture(page['background'],Pt(ox),Pt(oy),
                                  width=Pt(page['page_width']*scale),height=Pt(page['page_height']*scale))
    image.name='页面背景（已分离的文字、图片和图形已移除）'
    image.line.fill.background()
    for obj in graphic_objects(page):
        add_graphic(slide,obj,scale,ox,oy)
    from paragraph_layout import add_textbox
    for obj in text_objects(page):
        add_textbox(slide,obj,scale,ox,oy)
    description=('本页文字由本机 OCR 重建，同段文字已合并为可编辑文段；字体为相近匹配。'
                 if page.get('lines') else '本页经过版面分析与对象分离。')
    if page.get('images'):
        description+='已分离的图片可单独移动、缩放、裁剪或替换。'
    if page.get('shapes'):
        description+='已重建的简单图形可单独编辑。'
    slide.notes_slide.notes_text_frame.text=(description+
        '未可靠分离的内容仍保留在背景；识别结果和背景修复请校对。\n'+page.get('text',''))
    return slide


def image_pdf(src,dst):
    """Normalise oriented/multipage images without uploading or JPEG recompression."""
    import pymupdf as fitz
    from PIL import Image,ImageOps,ImageSequence
    with fitz.open() as document,Image.open(src) as original:
        for frame in ImageSequence.Iterator(original):
            image=ImageOps.exif_transpose(frame).convert('RGBA')
            factor=min(.75,2800/max(image.size))
            page=document.new_page(width=image.width*factor,height=image.height*factor)
            stream=io.BytesIO()
            image.save(stream,'PNG')
            page.insert_image(page.rect,stream=stream.getvalue())
            if Path(src).suffix.lower() not in {'.tif','.tiff'}:
                break
        document.save(dst)


def layout_with_ocr(src,dst,fmt='pptx',progress=None,force=False):
    from pdf_layout import layout_slides,layout_word
    with tempfile.TemporaryDirectory(prefix='.lightflip-ocr-',dir=Path(dst).resolve().parent) as temp:
        folder=Path(temp)
        progress=progress or (lambda message:None)
        if Path(src).suffix.lower()!='.pdf':
            normalised=folder/'images.pdf'
            image_pdf(src,normalised)
            src=normalised
        pages=prepare_pages(src,folder,progress,force)
        if fmt=='pptx':layout_slides(src,dst,ocr_pages=pages,progress=progress)
        else:layout_word(src,dst,ocr_pages=pages,progress=progress)


def extract_document_text(src,progress=None):
    """Use native PDF text when present, local OCR otherwise; images always OCR."""
    import pymupdf as fitz
    from file_tools import ai_job
    progress=progress or (lambda message:None)
    with tempfile.TemporaryDirectory(prefix='.lightflip-text-',dir=Path(src).resolve().parent) as temp:
        folder=Path(temp)
        if Path(src).suffix.lower()!='.pdf':
            normalized=folder/'images.pdf'
            image_pdf(src,normalized)
            src=normalized
        texts,sources,numbers=[],[],[]
        with fitz.open(src) as document:
            if document.needs_pass:raise ValueError('PDF 有密码保护，请先解锁。')
            for index,page in enumerate(document):
                text=page.get_text().strip()
                texts.append(text)
                if image_page(page):
                    factor=min(3,3200/max(page.rect.width,page.rect.height))
                    file=folder/f'page-{index+1}.png'
                    page.get_pixmap(matrix=fitz.Matrix(factor,factor),alpha=False).save(file)
                    sources.append(str(file))
                    numbers.append(index)
        if sources:
            result=ai_job({'kind':'ocr_text','files':sources},folder,progress)
            for index,item in zip(numbers,result['items']):
                # Preserve small native annotations that were too faint to OCR.
                recognised=item['text']
                import re
                normalize=lambda value:re.sub(r'\W','',value).lower()
                native=[line for line in texts[index].splitlines()
                        if normalize(line) and normalize(line) not in normalize(recognised)]
                texts[index]='\n'.join([recognised,*native]).strip()
        if not any(text.strip() for text in texts):
            raise ValueError('未识别到文字。请检查文件清晰度或选择保留页面外观。')
        return '\n\n'.join(f'===== 第 {index+1} 页 =====\n'+text for index,text in enumerate(texts))
